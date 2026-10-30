"""Excel export of the period dashboard (E6-T3).

Sheet « KPIs » holds the period KPIs as the dashboard cards show them, sheet
« Lots » the lots sold during the period (one row per engaged reservation).
Numbers stay numbers and dates real dates, so the file can be reworked in Excel.
"""

import re
import unicodedata
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AuthorizationError, NotFoundError
from app.domain.dashboard_period import Period, business_today
from app.domain.lot_workflow import STATUS_LABELS
from app.domain.schemas.dashboard import PeriodRange, PeriodStats
from app.infrastructure.database.repositories import ProjectRepository
from app.services.dashboard_period_service import DashboardPeriodService, SoldLot

DATE_FORMAT = "DD/MM/YYYY"
KPI_HEADER = ("Indicateur", "Période", "Période précédente", "Variation (%)")
LOTS_HEADER = (
    "N° lot",
    "Zone",
    "Type de lot",
    "Surface (m²)",
    "Client",
    "Commercial",
    "Date d'engagement",
    "Prix total convenu (MAD)",
    "Prix acte (MAD)",
    "Statut actuel",
)
KPI_WIDTHS = (34, 60, 60, 14)
LOTS_WIDTHS = (12, 16, 18, 14, 28, 24, 18, 24, 18, 26)

# (label, FlowKpis field, based on the status history)
FLOW_ROWS = (
    ("Lots vendus", "lots_vendus", True),
    ("CA (MAD)", "ca", True),
    ("Réservations créées", "reservations_creees", False),
    ("Taux de conversion (%)", "taux_conversion", True),
    ("Encaissements (MAD)", "encaissements", False),
)


def slugify(name: str) -> str:
    """ASCII, lower-case slug of a project name (``"projet"`` if nothing is left)."""
    ascii_name = (
        unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")
    return slug or "projet"


def export_filename(project_name: str, period: Period | PeriodRange) -> str:
    """``dashboard_<slug>_<date_from>_<date_to>.xlsx`` (ISO dates)."""
    return (
        f"dashboard_{slugify(project_name)}_"
        f"{period.date_from.isoformat()}_{period.date_to.isoformat()}.xlsx"
    )


def _fr(day: date) -> str:
    return day.strftime("%d/%m/%Y")


def _range_label(period: PeriodRange) -> str:
    return f"du {_fr(period.date_from)} au {_fr(period.date_to)}"


def _bold_row(sheet: Worksheet, row: int) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=True)


def _set_widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_kpis(sheet: Worksheet, project_name: str, stats: PeriodStats) -> None:
    history_start = _fr(stats.history_start) if stats.history_start else ""
    flow_notice = f"Ventes, CA et conversion disponibles à partir du {history_start}"
    state_notice = f"Situation des lots disponible à partir du {history_start}"
    # Same rule as the variations: previous flows exist from the history start on.
    previous_available = (
        stats.history_start is None
        or stats.previous_period.date_from >= stats.history_start
    )

    sheet.append(("Projet", project_name))
    sheet.append(("Période", _range_label(stats.period)))
    sheet.append(("Période précédente", _range_label(stats.previous_period)))
    sheet.append(())
    sheet.append(KPI_HEADER)
    _bold_row(sheet, sheet.max_row)

    for label, field, needs_history in FLOW_ROWS:
        current: Any = getattr(stats.flow, field)
        previous: Any = getattr(stats.flow_previous, field)
        change: Any = getattr(stats.variations, field)
        if needs_history and not stats.flow_kpis_available:
            current, change = flow_notice, None
        if needs_history and not previous_available:
            previous, change = flow_notice, None
        sheet.append((label, current, previous, change))

    state_date = _fr(stats.period.date_to)
    for label, value in (
        (f"Lots disponibles au {state_date}", stats.state.available),
        (f"Lots en option au {state_date}", stats.state.option),
    ):
        sheet.append((label, value if stats.state_kpis_available else state_notice))

    _set_widths(sheet, KPI_WIDTHS)


def _write_lots(sheet: Worksheet, sold: list[SoldLot]) -> None:
    sheet.append(LOTS_HEADER)
    _bold_row(sheet, 1)
    for row in sold:
        sheet.append(
            (
                row.lot_numero,
                row.zone,
                row.type_lot,
                row.surface,
                row.client_name,
                row.commercial_name,
                row.engaged_on,
                row.agreed_total_price,
                row.deed_price,
                STATUS_LABELS.get(row.lot_status, row.lot_status),
            )
        )
        sheet.cell(row=sheet.max_row, column=7).number_format = DATE_FORMAT
    _set_widths(sheet, LOTS_WIDTHS)


def build_dashboard_workbook(
    project_name: str, stats: PeriodStats, sold: list[SoldLot]
) -> bytes:
    """The ``.xlsx`` content: sheets « KPIs » and « Lots »."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])  # the default empty sheet
    _write_kpis(workbook.create_sheet("KPIs"), project_name, stats)
    _write_lots(workbook.create_sheet("Lots"), sold)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class DashboardExportService:
    """Builds the Excel export of a project's period dashboard."""

    def __init__(self, session: AsyncSession):
        self.session = session
        self.project_repo = ProjectRepository(session)
        self.period_service = DashboardPeriodService(session)

    async def export(
        self,
        project_id: int,
        period: Period,
        user_id: int | None,
        current_user_id: int,
        current_user_role: str,
    ) -> tuple[str, bytes]:
        """File name and content of the export (``user_id`` filters the flows).

        Raises:
            NotFoundError: If the project does not exist
            AuthorizationError: If the user cannot access the project
        """
        project = await self.project_repo.get_by_id(project_id)
        if project is None:
            raise NotFoundError("Project", project_id)
        if not await self.project_repo.can_user_access(
            project_id, current_user_id, current_user_role
        ):
            raise AuthorizationError(message="Vous n'avez pas accès à ce projet.")

        stats = PeriodStats(
            **await self.period_service.period_stats(
                period, project_id, user_id, business_today()
            )
        )
        sold = await self.period_service.sold_lots(period, project_id, user_id)
        return (
            export_filename(project.name, period),
            build_dashboard_workbook(project.name, stats, sold),
        )
