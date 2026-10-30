"""GET /api/dashboard/export: Excel export of the period dashboard (E6-T3)."""

from collections.abc import Callable
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.domain.dashboard_period import business_tz
from app.infrastructure.database.models import (
    ClientModel,
    ProjectModel,
    ReservationModel,
    UserModel,
)
from app.tests.factories import (
    assign,
    make_client,
    make_lot,
    make_project,
    make_reservation,
    make_reservation_payment,
    make_status_change,
)

pytestmark = pytest.mark.integration

Headers = Callable[[UserModel], dict[str, str]]

URL = "/api/dashboard/export"
OCTOBER = {"date_from": "2026-10-01", "date_to": "2026-10-31"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KPI_HEADER = ("Indicateur", "Période", "Période précédente", "Variation (%)")


def _local(year: int, month: int, day: int, hour: int = 12) -> datetime:
    """UTC instant of a Casablanca local time (never hard-code the offset)."""
    return datetime(year, month, day, hour, tzinfo=business_tz()).astimezone(UTC)


async def _engaged(
    db: AsyncSession,
    project: ProjectModel,
    client_: ClientModel,
    on: datetime,
    *,
    user: UserModel | None = None,
    agreed: float = 500000.0,
    **lot_fields: Any,
) -> ReservationModel:
    lot = await make_lot(db, project, status="reservation_engagee", **lot_fields)
    reservation = await make_reservation(
        db, lot, client_, user=user, reservation_date=on, agreed_total_price=agreed
    )
    await make_status_change(
        db, lot, "reservation_engagee", on, reservation=reservation
    )
    return reservation


async def _history_since_august(db: AsyncSession, project: ProjectModel) -> None:
    """History starting in August: every KPI of October is available."""
    early = await make_lot(db, project, status="option")
    await make_status_change(db, early, "option", _local(2026, 8, 1))


def _workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(content))


def _lot_numbers(workbook: openpyxl.Workbook) -> list[str]:
    rows = list(workbook["Lots"].iter_rows(min_row=2, values_only=True))
    return [row[0] for row in rows]


def _kpis(workbook: openpyxl.Workbook) -> dict[str, Any]:
    """Current-period value of each indicator of the « KPIs » sheet."""
    rows = [tuple(row) for row in workbook["KPIs"].iter_rows(values_only=True)]
    start = rows.index(KPI_HEADER) + 1
    return {row[0]: row[1] for row in rows[start:] if row[0] is not None}


async def test_manager_export(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    commercial: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager, name="Résidence Al Amal / T2")
    buyer = await make_client(db_session, manager, name="Alami")
    await _history_since_august(db_session, project)
    late = await _engaged(
        db_session,
        project,
        buyer,
        _local(2026, 10, 20),
        user=commercial,
        agreed=650000.0,
        numero="B-2",
    )
    await _engaged(db_session, project, buyer, _local(2026, 10, 3), numero="A-1")
    await _engaged(db_session, project, buyer, _local(2026, 9, 15), numero="S-9")
    await make_reservation_payment(db_session, late, 1000.0, date(2026, 10, 25))
    headers = auth_headers(manager)

    response = await client.get(
        URL, headers=headers, params={"project_id": project.id, **OCTOBER}
    )
    stats = await client.get(
        "/api/dashboard/stats",
        headers=headers,
        params={"project_id": project.id, **OCTOBER},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == (
        'attachment; filename="dashboard_residence-al-amal-t2_2026-10-01_2026-10-31.xlsx"'
    )
    workbook = _workbook(response.content)
    assert workbook.sheetnames == ["KPIs", "Lots"]
    assert _lot_numbers(workbook) == ["A-1", "B-2"]
    flow = stats.json()["flow"]
    kpis = _kpis(workbook)
    assert kpis["Lots vendus"] == flow["lots_vendus"] == 2
    assert kpis["CA (MAD)"] == flow["ca"] == 1150000.0
    assert kpis["Réservations créées"] == flow["reservations_creees"]
    assert kpis["Taux de conversion (%)"] == flow["taux_conversion"]
    assert kpis["Encaissements (MAD)"] == flow["encaissements"] == 1000.0
    assert kpis["Lots disponibles au 31/10/2026"] == stats.json()["state"]["available"]
    assert kpis["Lots en option au 31/10/2026"] == stats.json()["state"]["option"]


async def test_commercial_not_assigned_is_refused(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    commercial: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager)

    response = await client.get(
        URL,
        headers=auth_headers(commercial),
        params={"project_id": project.id, **OCTOBER},
    )

    assert response.status_code == 403, response.text
    assert response.json()["message"] == "Vous n'avez pas accès à ce projet."


async def test_commercial_exports_only_their_own_sales(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    commercial: UserModel,
    other_commercial: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager)
    await assign(db_session, commercial, project)
    buyer = await make_client(db_session, manager)
    await _history_since_august(db_session, project)
    await _engaged(
        db_session, project, buyer, _local(2026, 10, 5), user=commercial, numero="M-1"
    )
    await _engaged(
        db_session,
        project,
        buyer,
        _local(2026, 10, 6),
        user=other_commercial,
        numero="T-1",
    )

    response = await client.get(
        URL,
        headers=auth_headers(commercial),
        params={"project_id": project.id, "user_id": other_commercial.id, **OCTOBER},
    )

    assert response.status_code == 200, response.text
    workbook = _workbook(response.content)
    assert _lot_numbers(workbook) == ["M-1"]
    assert _kpis(workbook)["Lots vendus"] == 1


async def test_manager_user_filter(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    commercial: UserModel,
    other_commercial: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager)
    buyer = await make_client(db_session, manager)
    await _engaged(
        db_session, project, buyer, _local(2026, 10, 5), user=commercial, numero="M-1"
    )
    await _engaged(
        db_session,
        project,
        buyer,
        _local(2026, 10, 6),
        user=other_commercial,
        numero="T-1",
    )

    response = await client.get(
        URL,
        headers=auth_headers(manager),
        params={"project_id": project.id, "user_id": other_commercial.id, **OCTOBER},
    )

    assert response.status_code == 200, response.text
    assert _lot_numbers(_workbook(response.content)) == ["T-1"]


async def test_missing_project_id_returns_422(
    client: AsyncClient,
    manager: UserModel,
    auth_headers: Headers,
) -> None:
    response = await client.get(URL, headers=auth_headers(manager), params=OCTOBER)

    assert response.status_code == 422, response.text


async def test_unknown_project_returns_404(
    client: AsyncClient,
    manager: UserModel,
    auth_headers: Headers,
) -> None:
    response = await client.get(
        URL, headers=auth_headers(manager), params={"project_id": 9999, **OCTOBER}
    )

    assert response.status_code == 404, response.text


async def test_inverted_dates_return_422(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager)

    response = await client.get(
        URL,
        headers=auth_headers(manager),
        params={
            "project_id": project.id,
            "date_from": "2026-10-31",
            "date_to": "2026-10-01",
        },
    )

    assert response.status_code == 422, response.text
    assert (
        response.json()["message"]
        == "La date de début doit être antérieure ou égale à la date de fin."
    )


async def test_cors_exposes_content_disposition(
    client: AsyncClient,
    db_session: AsyncSession,
    manager: UserModel,
    auth_headers: Headers,
) -> None:
    project = await make_project(db_session, manager)

    response = await client.get(
        URL,
        headers={**auth_headers(manager), "Origin": settings.CORS_ORIGINS[0]},
        params={"project_id": project.id, **OCTOBER},
    )

    assert response.status_code == 200, response.text
    exposed = response.headers["access-control-expose-headers"]
    assert "Content-Disposition" in [name.strip() for name in exposed.split(",")]
