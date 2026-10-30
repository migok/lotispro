"""Excel export of the period dashboard (E6-T3): file name and workbook."""

from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import pytest

from app.domain.dashboard_period import Period
from app.domain.schemas.dashboard import (
    FlowKpis,
    FlowVariations,
    PeriodRange,
    PeriodStats,
    StateKpis,
)
from app.services.dashboard_export import (
    build_dashboard_workbook,
    export_filename,
    slugify,
)
from app.services.dashboard_period_service import SoldLot

OCTOBER = PeriodRange(date_from=date(2026, 10, 1), date_to=date(2026, 10, 31))
SEPTEMBER = PeriodRange(date_from=date(2026, 8, 31), date_to=date(2026, 9, 30))

KPI_HEADER = ("Indicateur", "Période", "Période précédente", "Variation (%)")
LOTS_HEADER = (
    "N° lot",
    "Zone",
    "Type de lot",
    "Surface (m²)",
    "Client",
    "Commercial",
    "Date d'engagement",
    "Prix total convenu (MAD)",
    "Prix acte (MAD)",
    "Statut actuel",
)
FLOW_NOTICE = "Ventes, CA et conversion disponibles à partir du 10/10/2026"
STATE_NOTICE = "Situation des lots disponible à partir du 10/10/2026"


def _stats(**overrides: Any) -> PeriodStats:
    fields: dict[str, Any] = {
        "period": OCTOBER,
        "previous_period": SEPTEMBER,
        "flow": FlowKpis(
            lots_vendus=3,
            ca=1500000.5,
            reservations_creees=4,
            taux_conversion=75.0,
            encaissements=20000.0,
        ),
        "flow_previous": FlowKpis(
            lots_vendus=2,
            ca=1000000.0,
            reservations_creees=4,
            taux_conversion=50.0,
            encaissements=0.0,
        ),
        "variations": FlowVariations(
            lots_vendus=50.0,
            ca=50.0,
            reservations_creees=0.0,
            taux_conversion=50.0,
            encaissements=None,
        ),
        "state": StateKpis(available=7, option=2),
        "state_kpis_available": True,
        "flow_kpis_available": True,
        "history_start": date(2026, 1, 1),
    }
    return PeriodStats(**{**fields, **overrides})


def _sold(**overrides: Any) -> SoldLot:
    fields: dict[str, Any] = {
        "reservation_id": 1,
        "lot_id": 10,
        "lot_numero": "A-1",
        "zone": "Nord",
        "type_lot": "Villa",
        "surface": 250.5,
        "client_name": "Alami",
        "commercial_name": "Karim",
        "engaged_on": date(2026, 10, 5),
        "agreed_total_price": 650000.0,
        "deed_price": 600000.0,
        "lot_status": "reservation_engagee",
    }
    return SoldLot(**{**fields, **overrides})


def _load(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(content))


def _rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def _kpi_rows(sheet: Any) -> dict[str, tuple[Any, ...]]:
    """Indicator rows of the « KPIs » sheet, keyed by label."""
    rows = _rows(sheet)
    start = rows.index(KPI_HEADER) + 1
    return {row[0]: row[1:] for row in rows[start:] if row[0] is not None}


@pytest.mark.parametrize(
    ("name", "slug"),
    [
        ("Résidence Al Amal / Tranche 2", "residence-al-amal-tranche-2"),
        ("  --Les Jardins d'Été--  ", "les-jardins-d-ete"),
        ("Projet   ÀÉÎ", "projet-aei"),
        ("مشروع", "projet"),
        ("", "projet"),
    ],
)
def test_slugify(name: str, slug: str) -> None:
    assert slugify(name) == slug


def test_export_filename() -> None:
    period = Period(date(2026, 10, 1), date(2026, 10, 31))

    assert (
        export_filename("Résidence Al Amal / Tranche 2", period)
        == "dashboard_residence-al-amal-tranche-2_2026-10-01_2026-10-31.xlsx"
    )


def test_workbook_sheets_and_kpis() -> None:
    workbook = _load(build_dashboard_workbook("Al Amal", _stats(), [_sold()]))

    assert workbook.sheetnames == ["KPIs", "Lots"]
    sheet = workbook["KPIs"]
    assert _rows(sheet)[:3] == [
        ("Projet", "Al Amal", None, None),
        ("Période", "du 01/10/2026 au 31/10/2026", None, None),
        ("Période précédente", "du 31/08/2026 au 30/09/2026", None, None),
    ]
    assert _kpi_rows(sheet) == {
        "Lots vendus": (3, 2, 50.0),
        "CA (MAD)": (1500000.5, 1000000.0, 50.0),
        "Réservations créées": (4, 4, 0.0),
        "Taux de conversion (%)": (75.0, 50.0, 50.0),
        "Encaissements (MAD)": (20000.0, 0.0, None),
        "Lots disponibles au 31/10/2026": (7, None, None),
        "Lots en option au 31/10/2026": (2, None, None),
    }
    header_row = _rows(sheet).index(KPI_HEADER) + 1
    assert all(cell.font.bold for cell in sheet[header_row])


def test_workbook_lots_sheet() -> None:
    sold = [
        _sold(),
        _sold(
            reservation_id=2,
            lot_numero="B-2",
            zone=None,
            type_lot=None,
            surface=None,
            client_name=None,
            commercial_name=None,
            engaged_on=date(2026, 10, 20),
            agreed_total_price=0.0,
            deed_price=None,
            lot_status="chez_proprietaire",
        ),
    ]

    sheet = _load(build_dashboard_workbook("Al Amal", _stats(), sold))["Lots"]

    rows = _rows(sheet)
    assert rows[0] == LOTS_HEADER
    assert all(cell.font.bold for cell in sheet[1])
    assert rows[1:] == [
        (
            "A-1",
            "Nord",
            "Villa",
            250.5,
            "Alami",
            "Karim",
            datetime(2026, 10, 5),
            650000.0,
            600000.0,
            "Réservation engagée",
        ),
        (
            "B-2",
            None,
            None,
            None,
            None,
            None,
            datetime(2026, 10, 20),
            0,
            None,
            "Chez le propriétaire",
        ),
    ]
    assert sheet["G2"].number_format == "DD/MM/YYYY"
    assert sheet["G2"].is_date


def test_workbook_without_sold_lots_has_only_the_header() -> None:
    sheet = _load(build_dashboard_workbook("Al Amal", _stats(), []))["Lots"]

    assert _rows(sheet) == [LOTS_HEADER]


def test_unavailable_values_use_the_ui_notices() -> None:
    stats = _stats(
        flow_kpis_available=False,
        state_kpis_available=False,
        variations=FlowVariations(),
        history_start=date(2026, 10, 10),
    )

    rows = _kpi_rows(_load(build_dashboard_workbook("Al Amal", stats, []))["KPIs"])

    assert rows["Lots vendus"] == (FLOW_NOTICE, FLOW_NOTICE, None)
    assert rows["CA (MAD)"] == (FLOW_NOTICE, FLOW_NOTICE, None)
    assert rows["Taux de conversion (%)"] == (FLOW_NOTICE, FLOW_NOTICE, None)
    # Not based on the status history: still exported.
    assert rows["Réservations créées"] == (4, 4, None)
    assert rows["Encaissements (MAD)"] == (20000.0, 0.0, None)
    assert rows["Lots disponibles au 31/10/2026"] == (STATE_NOTICE, None, None)
    assert rows["Lots en option au 31/10/2026"] == (STATE_NOTICE, None, None)


def test_previous_period_before_history_start_uses_the_notice() -> None:
    # Current flows available, previous period starts before the history.
    stats = _stats(
        period=PeriodRange(date_from=date(2026, 10, 10), date_to=date(2026, 10, 31)),
        previous_period=PeriodRange(
            date_from=date(2026, 9, 18), date_to=date(2026, 10, 9)
        ),
        variations=FlowVariations(),
        history_start=date(2026, 10, 10),
    )

    rows = _kpi_rows(_load(build_dashboard_workbook("Al Amal", stats, []))["KPIs"])

    assert rows["Lots vendus"] == (3, FLOW_NOTICE, None)
    assert rows["Réservations créées"] == (4, 4, None)
